"""Rebuild the List of Figures and List of Tables in the university report."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from pypdf import PdfReader


TABLE_CAPTION_FIXES = {
    "Table 3.1 Rooms Table": "Table 3.2 Rooms Table",
    "Table 3.2 Booking Table": "Table 3.3 Booking Table",
    "Table 3.3 Payments Table": "Table 3.4 Payments Table",
    "Table 3.4 Admin Table": "Table 3.5 Admin Table",
}


def normalize(text: str) -> str:
    """Normalize whitespace for matching DOCX captions against PDF text."""
    return re.sub(r"\s+", " ", text.replace("\u00a0", " ")).strip().casefold()


def find_paragraph(document: Document, text: str):
    """Find a paragraph by its trimmed visible text."""
    for paragraph in document.paragraphs:
        if paragraph.text.strip().casefold() == text.casefold():
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def remove_between(start_paragraph, end_paragraph) -> None:
    """Remove all body elements between two paragraph elements."""
    parent = start_paragraph._p.getparent()
    current = start_paragraph._p.getnext()
    while current is not None and current is not end_paragraph._p:
        following = current.getnext()
        parent.remove(current)
        current = following


def insert_entries_before(document: Document, anchor, entries: list[tuple[str, int]]) -> None:
    """Insert list entries before an anchor using the document's existing list style."""
    for caption, page_number in entries:
        paragraph = document.add_paragraph(style="table of figures")
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1
        paragraph.add_run(caption)
        paragraph.add_run(f"\t{page_number}")
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
        anchor._p.addprevious(paragraph._p)


def collect_captions(document: Document) -> tuple[list[str], list[str]]:
    """Collect figure and table captions from the report body in document order."""
    chapter_started = False
    figures: list[str] = []
    tables: list[str] = []

    # Include captions placed inside table cells, which Document.paragraphs omits.
    all_paragraphs = [
        Paragraph(element, document._body)
        for element in document.element.body.iter(qn("w:p"))
    ]

    for paragraph in all_paragraphs:
        text = paragraph.text.strip()
        if text.upper() == "CHAPTER ONE":
            chapter_started = True
        if not chapter_started:
            continue

        if text in TABLE_CAPTION_FIXES:
            paragraph.text = TABLE_CAPTION_FIXES[text]
            paragraph.style = "Caption"
            text = paragraph.text.strip()

        if re.match(r"^(?:Appendix\s+)?Figure\s+", text, flags=re.I):
            figures.append(text)
        elif re.match(r"^Table\s+", text, flags=re.I):
            tables.append(text)

    return figures, tables


def locate_pages(pdf_path: Path, captions: list[str]) -> dict[str, int]:
    """Locate each caption on the rendered PDF, ignoring its front-matter list entry."""
    reader = PdfReader(pdf_path)
    page_texts = [normalize(page.extract_text() or "") for page in reader.pages]
    page_map: dict[str, int] = {}

    for caption in captions:
        target = normalize(caption)
        matches = [
            page_number
            for page_number, page_text in enumerate(page_texts, start=1)
            if target in page_text
        ]
        if not matches:
            raise ValueError(f"Caption not found in rendered PDF: {caption}")
        page_map[caption] = max(matches)

    return page_map


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_docx", type=Path)
    parser.add_argument("output_docx", type=Path)
    parser.add_argument("--pdf", type=Path)
    args = parser.parse_args()

    document = Document(args.input_docx)
    figures, tables = collect_captions(document)

    if args.pdf:
        page_map = locate_pages(args.pdf, figures + tables)
    else:
        page_map = {caption: 0 for caption in figures + tables}

    figure_heading = find_paragraph(document, "LIST OF FIGURES")
    table_heading = find_paragraph(document, "LIST OF TABLES")
    chapter_one = find_paragraph(document, "CHAPTER ONE")

    remove_between(figure_heading, table_heading)
    remove_between(table_heading, chapter_one)

    insert_entries_before(
        document,
        table_heading,
        [(caption, page_map[caption]) for caption in figures],
    )
    insert_entries_before(
        document,
        chapter_one,
        [(caption, page_map[caption]) for caption in tables],
    )

    # Give every preliminary list and the first chapter a clean page start.
    figure_heading.paragraph_format.page_break_before = True
    figure_heading.paragraph_format.keep_with_next = True
    table_heading.paragraph_format.page_break_before = True
    table_heading.paragraph_format.keep_with_next = True
    for heading in (figure_heading, table_heading):
        for run in heading.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)
    chapter_one.paragraph_format.page_break_before = True

    args.output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output_docx)

    print(f"figures={len(figures)}")
    print(f"tables={len(tables)}")
    print(f"output={args.output_docx}")


if __name__ == "__main__":
    main()
