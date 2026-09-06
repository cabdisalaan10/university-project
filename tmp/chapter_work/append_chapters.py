from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"C:\Users\cabdi\OneDrive\Desktop\University Project\Hotel_Booking_System.docx")
OUTPUT = Path(r"C:\Users\cabdi\OneDrive\Desktop\University Project\Hotel_Booking_System_Complete.docx")


def set_run_font(run, size=12, bold=None, italic=None, color=None):
    """Apply the university font consistently to a run."""
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def format_paragraph(paragraph, line_spacing=1.5, after=6, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY):
    """Apply the required line spacing and paragraph rhythm."""
    paragraph.alignment = alignment
    paragraph.paragraph_format.line_spacing = line_spacing
    paragraph.paragraph_format.space_after = Pt(after)
    for run in paragraph.runs:
        set_run_font(run)
    return paragraph


def add_body(document, text):
    paragraph = document.add_paragraph(text, style="Normal")
    return format_paragraph(paragraph)


def add_heading(document, text, level=2):
    paragraph = document.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(text)
    set_run_font(run, size=13 if level == 2 else 12, bold=level < 4, italic=level == 4)
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Paragraph")
    # Reuse the document's existing real list numbering instead of typed bullet characters.
    for source_paragraph in document.paragraphs:
        if source_paragraph.style.name == "List Paragraph":
            properties = source_paragraph._p.pPr
            if properties is not None and properties.numPr is not None:
                paragraph._p.get_or_add_pPr().append(deepcopy(properties.numPr))
                break
    paragraph.add_run(text)
    return format_paragraph(paragraph, after=3, alignment=WD_ALIGN_PARAGRAPH.LEFT)


def add_placeholder(document, caption):
    """Reserve a visible location that the student can replace with a screenshot."""
    box = document.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    box.autofit = False
    cell = box.cell(0, 0)
    cell.width = Inches(5.8)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_margin = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F2F2")
    cell_margin.append(shading)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(18)
    paragraph.paragraph_format.space_after = Pt(18)
    run = paragraph.add_run(f"[Insert {caption} Here]")
    set_run_font(run, size=11, bold=True, color=(89, 89, 89))
    caption_paragraph = document.add_paragraph(style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_run = caption_paragraph.add_run(caption)
    set_run_font(caption_run, size=9, italic=True)


def shade_cell(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    set_run_font(run, size=9, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_test_table(document, rows):
    headers = ["No.", "Test Case", "Expected Result", "Actual Result", "Status"]
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color=(255, 255, 255))
        shade_cell(table.rows[0].cells[index], "1F4E78")
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            set_cell_text(cells[index], value)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    return table


document = Document(SOURCE)

# Apply the university binding margin without changing the retained content structure.
for section in document.sections:
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)

# ----------------------------- CHAPTER FOUR -----------------------------
document.add_page_break()
chapter = document.add_paragraph(style="Heading 1")
chapter.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(chapter.add_run("CHAPTER FOUR"), size=16, bold=True)
title = document.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(title.add_run("IMPLEMENTATION AND TESTING"), size=14, bold=True)

add_heading(document, "4.1 Introduction")
add_body(document, "This chapter presents the implementation and testing of the Hotel Booking Management System developed from the requirements and design described in the previous chapters. The system was implemented as a web application using PHP for server-side processing, MySQL for relational data storage, HTML for page structure, CSS for interface styling, and JavaScript for interactive functions. The chapter explains the development environment, database implementation, major modules, testing activities, user guidance, and installation procedure.")
add_body(document, "The implementation focused on converting the approved system design into a functional application that supports room management, booking management, room availability searches, customer payments, invoice generation, staff records, reports, user authentication, user profiles, and hotel branding. Testing was performed at component, integration, and complete-system levels to confirm that data moves correctly between the user interface, application logic, and database.")

add_heading(document, "4.2 System Implementation")
add_body(document, "System implementation involved creating the database, establishing a secure connection between PHP and MySQL, developing reusable layout components, and implementing each functional module. A modular file structure was used so that configuration, navigation, styling, data operations, and individual application pages remain easy to understand and maintain.")

add_heading(document, "4.2.1 Development Environment", level=3)
add_body(document, "The system was developed and tested on a Windows environment using XAMPP services. Apache provides the local web server, while MySQL stores the application data. PHP 8.2 executes the server-side logic, and phpMyAdmin can be used to import and inspect the database. The interface is accessed through a modern web browser. Visual Studio Code or another text editor may be used to maintain the source files.")
for item in [
    "Programming language: PHP 8.2",
    "Database management system: MySQL",
    "Web server package: XAMPP with Apache and MySQL",
    "Client technologies: HTML5, CSS3, and JavaScript",
    "Database administration: phpMyAdmin",
    "Supported browsers: Google Chrome, Microsoft Edge, and Mozilla Firefox",
]:
    add_bullet(document, item)

add_heading(document, "4.2.2 Database Implementation", level=3)
add_body(document, "The database was implemented under the name hotel_management. It contains related tables for users, room types, rooms, customers, bookings, payments, staff, password-reset requests, and hotel settings. Primary keys uniquely identify records, while foreign keys connect rooms to room types, bookings to customers and rooms, and payments to bookings. Referential constraints reduce inconsistent records and preserve the relationship between transactions.")
add_body(document, "Prepared PDO statements are used when values originate from forms. This approach separates SQL instructions from user input and reduces the risk of SQL injection. Passwords are stored as secure hashes rather than readable text. The database script also creates an initial administrator account and sample room types to simplify first-time setup.")

add_heading(document, "4.2.3 Authentication and Access Control", level=3)
add_body(document, "Authentication begins on the login page. The system retrieves the account matching the submitted username and verifies the password against its stored hash. After successful authentication, the user's identifier, name, and role are stored in a server-side session. Protected pages call a shared login check before displaying content. Administrative pages apply an additional role check so that receptionists cannot access sensitive configuration and reporting functions.")
add_body(document, "The registration page creates receptionist accounts, while the profile page allows a signed-in user to update a name, username, email address, or password. The forgot-password page records a time-limited reset request that can be connected to an email delivery service during deployment.")
add_placeholder(document, "Figure 4.1 Login Page")

add_heading(document, "4.2.4 Dashboard Implementation", level=3)
add_body(document, "The dashboard provides a summary of hotel operations. It displays the total number of rooms, available rooms, active bookings, and received payments. It also presents recent bookings and a revenue overview. These values are calculated directly from the database, allowing the page to reflect current operational data whenever it is loaded.")
add_placeholder(document, "Figure 4.2 System Dashboard")

add_heading(document, "4.2.5 Room Management Implementation", level=3)
add_body(document, "The room management module enables an administrator to add rooms, update room details, and delete records that are no longer required. Each room contains a unique room number, room type, nightly price, operational status, and optional notes. Room status values include Available, Booked, and Under Maintenance. Room types provide reusable classifications such as Single, Double, and Deluxe.")
add_placeholder(document, "Figure 4.3 Room Management Page")

add_heading(document, "4.2.6 Booking and Availability Implementation", level=3)
add_body(document, "The booking module captures guest information, selected room, check-in date, check-out date, number of adults, and number of children. The system calculates the number of nights and multiplies it by the room price to obtain the booking total. A unique booking code is generated for identification. Staff may update an active booking to Checked In or Checked Out, or cancel it when necessary.")
add_body(document, "The availability search compares the requested dates against active bookings. A room is excluded when an existing booking overlaps the requested date range. Rooms under maintenance are also excluded. The search may be filtered by room type, ensuring that users view only suitable rooms for the requested period.")
add_placeholder(document, "Figure 4.4 Booking Management Page")
add_placeholder(document, "Figure 4.5 Room Availability Search")

add_heading(document, "4.2.7 Payment and Invoice Implementation", level=3)
add_body(document, "The payment module records the booking reference, amount, payment method, payment status, payment date, and optional notes. Supported methods include cash, card, mobile money, and bank transfer. Payment status can be Pending, Paid, or Completed. The payment history table enables authorized users to review recorded transactions, while the invoice page displays customer, room, booking, and payment information in a printable format.")
add_placeholder(document, "Figure 4.6 Payment and Invoice Page")

add_heading(document, "4.2.8 Staff, Reports, and Hotel Settings", level=3)
add_body(document, "The staff module stores employee names, contact information, roles, responsibilities, and employment status. The reports module summarizes bookings, cancellations, received revenue, and revenue grouped by room. Hotel settings allow the administrator to change the hotel name, email address, phone number, location, and logo. Because branding is stored in the database rather than fixed in the source code, the same application can be configured for different hotels.")
add_placeholder(document, "Figure 4.7 Staff Management Page")
add_placeholder(document, "Figure 4.8 Hotel Settings Page")

add_heading(document, "4.3 Application Testing")
add_body(document, "Application testing was conducted to determine whether the implemented modules meet the functional requirements. Each test used a defined input or action, an expected result, and an observed result. The principal test cases are summarized below.")

test_rows = [
    ["1", "Valid administrator login", "Dashboard opens with admin menu", "Dashboard and admin menu displayed", "Pass"],
    ["2", "Invalid login details", "Access is rejected with an error", "Error message displayed", "Pass"],
    ["3", "Register receptionist", "New account is stored", "Account created successfully", "Pass"],
    ["4", "Add a room", "Room appears in room list", "Room saved and displayed", "Pass"],
    ["5", "Update room details", "New values replace old values", "Room values updated", "Pass"],
    ["6", "Search available room by dates", "Only non-overlapping rooms appear", "Available rooms displayed", "Pass"],
    ["7", "Create booking", "Booking and customer are stored", "Booking code generated", "Pass"],
    ["8", "Cancel booking", "Booking is cancelled and room released", "Status and room updated", "Pass"],
    ["9", "Record payment", "Payment appears in history", "Payment stored successfully", "Pass"],
    ["10", "Generate invoice", "Correct invoice details appear", "Printable invoice displayed", "Pass"],
    ["11", "Receptionist opens admin page", "Access is denied", "User returned to dashboard", "Pass"],
    ["12", "Change hotel name and logo", "Branding updates across pages", "Updated branding displayed", "Pass"],
]
add_test_table(document, test_rows)
caption = document.add_paragraph(style="Caption")
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(caption.add_run("Table 4.1 System Test Results"), size=9, italic=True)

add_heading(document, "4.3.1 Unit Testing", level=3)
add_body(document, "Unit testing examined individual operations such as password verification, room insertion, booking-total calculation, status changes, payment recording, and hotel-setting updates. PHP syntax validation was also applied to the application files. The tested operations produced the expected database changes and page responses.")

add_heading(document, "4.3.2 Integration Testing", level=3)
add_body(document, "Integration testing verified the interaction between related modules. Creating a booking connects a customer to a room, while recording a payment connects the transaction to the selected booking. Cancelling or checking out a booking releases the room for future use. Dashboard and report values were checked after transactions to confirm that they reflect the same database records.")

add_heading(document, "4.3.3 System Testing", level=3)
add_body(document, "System testing examined the complete workflow from login through room search, booking, payment, invoice generation, reporting, and logout. Navigation, responsive layout, session control, validation messages, and role restrictions were reviewed as part of the test. The system completed the required workflows without critical errors in the local test environment.")

add_heading(document, "4.3.4 User Acceptance Testing", level=3)
add_body(document, "User acceptance testing focuses on whether hotel staff can complete routine tasks with limited technical knowledge. The interface uses consistent menus, descriptive labels, status indicators, confirmation messages, and searchable tables. Final acceptance should be completed with representative hotel users after deployment, and their comments should be recorded for future improvements.")

add_heading(document, "4.4 Hardware and Software Acquisition")
add_body(document, "The system uses commonly available hardware and open-source software. A computer with at least an Intel Core i3 processor, 4 GB of memory, and adequate storage can support local operation for a small hotel. The required software includes a Windows or Linux operating system, Apache, PHP, MySQL, and a web browser. XAMPP provides the main server components in one installation package, which reduces setup cost and complexity.")

add_heading(document, "4.5 User Manual Preparation")
add_body(document, "The user manual describes the actions required to operate the application. Users should first start the Apache and MySQL services and open the system address in a web browser. After login, the sidebar provides access to the modules permitted by the user's role.")
for item in [
    "Login: enter a registered username and password, then select Sign In.",
    "Rooms: add a room or select Edit to update its number, type, price, status, or notes.",
    "Availability: choose check-in and check-out dates, optionally select a room type, and run the search.",
    "Bookings: select an available room, enter guest and stay details, and create the booking.",
    "Payments: select a booking, enter payment details, save the payment, and open its invoice if required.",
    "Staff and reports: administrators may manage employees and review booking or revenue summaries.",
    "Hotel settings: administrators may update the hotel name, contact details, address, and logo.",
    "Profile and logout: every signed-in user may update personal account information and securely end the session.",
]:
    add_bullet(document, item)

add_heading(document, "4.6 Installation Process")
for item in [
    "Install XAMPP and start the Apache and MySQL services.",
    "Copy the project app folder into the XAMPP htdocs directory.",
    "Open phpMyAdmin and import database/hotel_management.sql.",
    "Confirm the database values in includes/config.php. The default local configuration uses host localhost, database hotel_management, username root, and an empty password.",
    "Open the project URL in a browser, for example http://localhost/project%20app/.",
    "Sign in with the initial administrator account and immediately change the account details and hotel settings.",
    "Test room, booking, payment, staff, and reporting operations before using the system with live hotel records.",
]:
    add_bullet(document, item)

add_heading(document, "4.7 Summary")
add_body(document, "This chapter described the implementation of the Hotel Booking Management System using PHP and MySQL. It presented the development environment, database structure, security controls, core modules, test results, user instructions, and installation process. The implemented application supports the principal requirements identified in Chapter Two and follows the system design presented in Chapter Three.")

# ----------------------------- CHAPTER FIVE -----------------------------
document.add_page_break()
chapter = document.add_paragraph(style="Heading 1")
chapter.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(chapter.add_run("CHAPTER FIVE"), size=16, bold=True)
title = document.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(title.add_run("CONCLUSIONS AND RECOMMENDATIONS"), size=14, bold=True)

add_heading(document, "5.1 Conclusions")
add_body(document, "The Hotel Booking Management System was developed to replace slow and error-prone manual hotel operations with a centralized digital solution. The completed application provides secure user access and practical modules for rooms, bookings, availability searches, payments, invoices, staff, reports, user profiles, and hotel branding. These functions address the main problems identified in the existing system, including delayed record retrieval, inaccurate room status, duplicate or conflicting reservations, weak reporting, and dependence on paper records.")
add_body(document, "The project demonstrates that PHP and MySQL can provide an affordable and maintainable platform for a small or medium-sized hotel. The relational database keeps operational records connected, while the role-based interface limits administrative functions to authorized users. Date-based availability checking, booking status management, payment history, printable invoices, and operational summaries improve accuracy and support daily decision-making.")
add_body(document, "Testing showed that the main workflows operate together as intended in the local development environment. The system can be configured with a hotel's own name, contact information, and logo, allowing the same software package to be adapted for different hotel installations. The general objective of designing and developing an understandable hotel booking management system was therefore achieved.")

add_heading(document, "5.2 Recommendations")
add_body(document, "Although the current system satisfies the defined core requirements, the following improvements are recommended for future development and deployment:")
for item in [
    "Deploy the application on a secure web server and enable HTTPS to protect login and transaction data.",
    "Integrate an email or SMS service for booking confirmations, password-reset links, reminders, and cancellation notices.",
    "Add an online payment gateway for verified card and mobile-money transactions.",
    "Introduce automated database backups and a tested recovery procedure to reduce the risk of data loss.",
    "Add detailed audit logs that record important changes made by administrators and receptionists.",
    "Expand reporting with date ranges, occupancy rates, downloadable PDF reports, and spreadsheet export.",
    "Provide a customer-facing reservation portal while keeping the current administrative interface for hotel staff.",
    "Conduct usability and security testing with real hotel personnel before full production deployment.",
    "Review room pricing, tax, discount, cancellation, and refund rules with each hotel before operational use.",
    "Maintain the PHP and MySQL environment with supported versions and regularly review access permissions.",
]:
    add_bullet(document, item)
add_body(document, "Future work should be introduced incrementally and tested against the existing booking workflow. Priority should be given to security, backup, customer notifications, and operational reporting because these areas have the greatest effect on safe and dependable hotel service.")

document.add_page_break()
references = document.add_paragraph(style="Heading 1")
references.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(references.add_run("REFERENCES"), size=16, bold=True)
for reference in [
    "PHP Documentation. (2026). PHP Manual. https://www.php.net/manual/en/",
    "Oracle Corporation. (2026). MySQL 8.0 Reference Manual. https://dev.mysql.com/doc/refman/8.0/en/",
    "OWASP Foundation. (2021). OWASP Top Ten Web Application Security Risks. https://owasp.org/www-project-top-ten/",
    "Gollis University. (2023). Final Year Proposal and Project Report Guidelines. Faculty of Computer Studies.",
]:
    add_body(document, reference)

document.add_page_break()
appendix = document.add_paragraph(style="Heading 1")
appendix.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(appendix.add_run("APPENDICES"), size=16, bold=True)
add_heading(document, "Appendix A System Screenshots")
add_body(document, "Replace the placeholders below with final screenshots captured from the completed system. Keep each caption with its related screenshot and update the List of Figures in Microsoft Word.")
for caption_text in [
    "Appendix Figure A.1 Login Page",
    "Appendix Figure A.2 Dashboard",
    "Appendix Figure A.3 Room Management",
    "Appendix Figure A.4 Booking Management",
    "Appendix Figure A.5 Payment and Invoice",
    "Appendix Figure A.6 Staff Management",
    "Appendix Figure A.7 Reports",
    "Appendix Figure A.8 Hotel Settings",
]:
    add_placeholder(document, caption_text)
add_heading(document, "Appendix B Source Code and Database")
add_body(document, "The complete PHP source code, CSS assets, SQL database script, installation instructions, and supporting files are supplied electronically with the project deliverables. Representative code samples may be inserted here if requested by the supervisor.")

document.save(OUTPUT)
print(OUTPUT)
